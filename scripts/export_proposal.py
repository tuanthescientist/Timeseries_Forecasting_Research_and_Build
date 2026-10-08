"""Export docs/proposal.md; never edit the generated Word copy as a second source."""
from __future__ import annotations

import re

from btcforecast.data import ROOT


def main():
    from docx import Document
    from docx.shared import Inches, Pt

    source = ROOT / "docs/proposal.md"
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    for block in re.split(r"\n\s*\n", source.read_text(encoding="utf-8").strip()):
        line = " ".join(block.splitlines())
        match = re.match(r"^(#{1,3}) (.*)", line)
        if match:
            document.add_heading(match[2], level=len(match[1]) - 1)
            continue
        # Keep source URLs visible in the export; substantive prose comes from Markdown.
        line = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", line)
        document.add_paragraph(line.replace("**", ""))
    output = ROOT / "results/local/proposal.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Exported {output.relative_to(ROOT)}; render and visually inspect before sharing.")


if __name__ == "__main__":
    main()
